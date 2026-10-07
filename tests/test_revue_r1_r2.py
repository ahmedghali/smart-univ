"""Tests pour la Phase 2 bis - Groupe 1 : R1 et R2."""

import io

import openpyxl
import pytest
import tablib
from django.contrib.messages.storage.fallback import FallbackStorage
from django.test import RequestFactory
from django.urls import reverse

from apps.academique.departement.dep_admin.admins import EtudiantDepAdmin
from apps.academique.departement.dep_admin.resources import (
    EnseignantResource,
    EtudiantResource,
    generate_credentials_excel,
)
from apps.academique.departement.dep_admin.site import dep_admin_site
from apps.academique.enseignant.models import Enseignant
from apps.academique.etudiant.models import Etudiant
from apps.noyau.authentification.models import CustomUser
from tests.factories import EnseignantFactory, EtudiantFactory


@pytest.mark.django_db
class TestR1CredentialsExcel:
    """R1: Les comptes créés par import ont doit_changer_mot_de_passe=True

    et un fichier Excel en mémoire est généré avec leurs identifiants (téléchargement unique).
    """

    def test_r1_generate_credentials_excel(self):
        """Vérifie le format et les colonnes du fichier Excel généré en mémoire."""
        accounts = [
            {"nom": "Benali", "prenom": "Mohamed", "username": "mohamed.benali", "password": "TempPassword123!"},
            {"nom": "Kaci", "prenom": "Fatima", "username": "fatima.kaci", "password": "TempPassword456!"},
        ]
        excel_bytes = generate_credentials_excel(accounts)
        assert isinstance(excel_bytes, bytes)
        assert len(excel_bytes) > 0

        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        ws = wb.active
        assert ws.title == "Identifiants"
        headers = [cell.value for cell in ws[1]]
        assert headers == ["Nom", "Prénom", "Identifiant", "Mot de passe temporaire"]

        row2 = [cell.value for cell in ws[2]]
        assert row2 == ["Benali", "Mohamed", "mohamed.benali", "TempPassword123!"]
        row3 = [cell.value for cell in ws[3]]
        assert row3 == ["Kaci", "Fatima", "fatima.kaci", "TempPassword456!"]

    def test_r1_enseignant_import_creates_user_with_flag_and_credentials(self):
        """Vérifie que l'import d'un enseignant crée un utilisateur avec doit_changer_mot_de_passe=True

        et remplit created_accounts.
        """
        resource = EnseignantResource()
        dataset = tablib.Dataset(
            ["Kamel", "كمال", "Cherif", "شريف", "ENS20269999"],
            headers=["prenom_fr", "prenom_ar", "nom_fr", "nom_ar", "matricule"],
        )
        result = resource.import_data(dataset, dry_run=False)

        assert not result.has_errors()
        enseignant = Enseignant.objects.get(matricule="ENS20269999")
        assert enseignant.user is not None
        assert enseignant.user.doit_changer_mot_de_passe is True

        assert hasattr(result, "created_accounts")
        assert len(result.created_accounts) == 1
        account = result.created_accounts[0]
        assert account["username"] == enseignant.user.username
        assert account["password"] != ""
        assert account["nom"] == "Cherif"
        assert account["prenom"] == "Kamel"

    def test_r1_etudiant_import_creates_user_with_flag_and_credentials(self):
        """Vérifie que l'import d'un étudiant crée un utilisateur avec doit_changer_mot_de_passe=True

        et transmet created_accounts via after_import.
        """
        etudiant = EtudiantFactory(user=None, matricule="ETU20268888", nom_fr="Dahmani", prenom_fr="Amine")
        resource = EtudiantResource()
        resource._imported_etudiant_ids = [etudiant.id]

        dataset = tablib.Dataset(headers=["matricule", "nom_fr", "prenom_fr"])
        # Mocking import result container
        from import_export.results import Result

        res = Result()
        resource.after_import(dataset, res, dry_run=False)

        etudiant.refresh_from_db()
        assert etudiant.user is not None
        assert etudiant.user.doit_changer_mot_de_passe is True
        assert hasattr(res, "created_accounts")
        assert len(res.created_accounts) == 1
        acc = res.created_accounts[0]
        assert acc["username"] == etudiant.user.username
        assert acc["password"] != ""
        assert acc["nom"] == "Dahmani"

    def test_r1_download_credentials_one_time_only(self):
        """Vérifie que la vue download-credentials ne permet qu'un seul téléchargement."""
        admin_instance = EtudiantDepAdmin(Etudiant, dep_admin_site)
        factory = RequestFactory()

        request = factory.get("/departement/admin/etudiant/etudiant/download-credentials/")
        request.user = CustomUser.objects.create_superuser(username="admin_r1", email="a@u.dz", password="p")
        # Ajout du support session et messages
        request.session = {
            "_import_credentials_etudiant": [
                {"nom": "Ali", "prenom": "Omar", "username": "omar.ali", "password": "Secret123!"}
            ]
        }
        request._messages = FallbackStorage(request)

        # Premier téléchargement -> Succès (200 avec fichier Excel)
        response1 = admin_instance.download_credentials_view(request)
        assert response1.status_code == 200
        assert response1["Content-Type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        assert 'filename="comptes_crees.xlsx"' in response1["Content-Disposition"]
        wb = openpyxl.load_workbook(io.BytesIO(response1.content))
        assert len(wb.sheetnames) >= 1

        # Deuxième téléchargement -> Les données de session ont été vidées (pop) -> Redirection
        response2 = admin_instance.download_credentials_view(request)
        assert response2.status_code == 302
        assert "_import_credentials_etudiant" not in request.session


@pytest.mark.django_db
class TestR2MustChangePasswordMiddleware:
    """R2: Chemins autorisés construits dynamiquement avec reverse(),

    test de bout en bout de la redirection et de la déconnexion.
    """

    def test_r2_etudiant_with_flag_redirected_and_can_logout(self, client):
        """Un étudiant avec le drapeau doit être redirigé vers son changement de mot de passe,

        mais peut toujours accéder à /logout/ sans être bloqué.
        """
        etudiant = EtudiantFactory()
        user = etudiant.user
        user.doit_changer_mot_de_passe = True
        user.save()

        client.force_login(user)

        # 1. Tentative d'accès à une page protégée -> Redirigé vers change password
        resp = client.get(reverse("etudiant:dashboard_Etud"))
        assert resp.status_code == 302
        assert reverse("etud:changePassword_Etud") in resp.url

        # 2. Accès à la déconnexion /logout/ -> Autorisé (GET affiche la confirmation 200, POST déconnecte 302)
        logout_url = reverse("auth:logout")
        assert logout_url == "/logout/"
        resp_logout_get = client.get(logout_url)
        assert resp_logout_get.status_code == 200

        resp_logout_post = client.post(logout_url)
        assert resp_logout_post.status_code == 302

        # L'utilisateur est déconnecté
        resp_after = client.get(reverse("etudiant:dashboard_Etud"))
        assert resp_after.status_code == 302
        assert "/login/" in resp_after.url

    def test_r2_enseignant_with_flag_redirected_and_allowed_routes(self, client):
        """Un enseignant avec le drapeau est redirigé vers sa page simple,

        et a accès à ses URLs de changement de mot de passe et /logout/.
        """
        enseignant = EnseignantFactory()
        user = enseignant.user
        user.doit_changer_mot_de_passe = True
        user.save()

        client.force_login(user)

        # 1. Tentative d'accès au dashboard général
        resp = client.get("/enseignant/1/dashboard/")
        assert resp.status_code == 302
        assert reverse("ense:change_password_Ens_simple") in resp.url

        # 2. Accès à /enseignant/change-password/ -> Autorisé
        resp_pw = client.get(reverse("ense:change_password_Ens_simple"))
        assert resp_pw.status_code == 200

        # 3. Accès à /logout/ -> Autorisé
        resp_logout = client.get(reverse("auth:logout"))
        assert resp_logout.status_code == 200
        resp_post = client.post(reverse("auth:logout"))
        assert resp_post.status_code == 302

    def test_r2_end_to_end_password_change_clears_flag(self, client):
        """Bout en bout : l'utilisateur change son mot de passe, le drapeau repasse à False,

        et il n'est plus redirigé.
        """
        etudiant = EtudiantFactory()
        user = etudiant.user
        user.set_password("OldTempSecret123!")
        user.doit_changer_mot_de_passe = True
        user.save()

        client.force_login(user)

        # Accès avant changement -> redirection
        resp = client.get(reverse("etudiant:dashboard_Etud"))
        assert resp.status_code == 302

        # Changement effectif de mot de passe
        resp_change = client.post(
            reverse("etud:changePassword_Etud"),
            {
                "old_password": "OldTempSecret123!",
                "new_password": "BrandNewSecret456!",
                "confirm_password": "BrandNewSecret456!",
            },
            follow=True,
        )
        assert resp_change.status_code == 200

        user.refresh_from_db()
        assert user.doit_changer_mot_de_passe is False

        # Accès après changement -> plus de redirection forcée par le middleware
        resp_post = client.get(reverse("etudiant:dashboard_Etud"))
        assert resp_post.status_code == 200
